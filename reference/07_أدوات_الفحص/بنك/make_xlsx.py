import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
bank=json.load(open('bank_p0b.json',encoding='utf-8')); meta=json.load(open('meta.json',encoding='utf-8'))
ar=json.load(open('audit_rows.json',encoding='utf-8'))
def nk(s): return re.sub(r'\s+',' ',re.sub(r'[؟?.,،:؛!«»"()\-–—*✓⟵]',' ',re.sub(r'[ً-ْـ]','',s))).strip()
# ملاحظات التدقيق المتوسطة والمنخفضة مربوطة بالبنود بنص السؤال
notes={}; unmatched=[]
for sev,c in ar['a1']+ar['a2']:
    cid=c[0][:4]; txt=c[1]
    if txt.startswith('CONCEPT_'): continue
    q=re.split(r'⟵| — ',txt)[0].strip().strip('«»').strip()
    q=re.sub(r'^\[[^\]]*\]\s*','',q).strip('«» ')
    hit=[x for x in bank.get(cid,[]) if nk(q) and (nk(q) in nk(x['q']) or nk(x['q']) in nk(q))]
    note=f"[{sev}] {c[2]} ← المقترح: {c[3]}"
    if len(hit)>=1:
        for h in hit[:1]: notes.setdefault(h['id'],[]).append(note)
    else: unmatched.append((cid,q[:50]))
c2u={}; 
for u in meta['units']:
    for c in meta['uc'].get(u,[]): c2u.setdefault(c,u)
stage_names={'1':'م1 (6–7)','2':'م2 (8–9)','3':'م3 (10–11)','4':'م4 (12–14)'}
FONT='Arial'; hdr_fill=PatternFill('solid',fgColor='2A7F8F'); in_fill=PatternFill('solid',fgColor='FFF6D5'); note_fill=PatternFill('solid',fgColor='FCE4E4')
thin=Side(style='thin',color='CCCCCC'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook()
# ورقة الإرشاد
L=wb.active; L.title='اقرأني'; L.sheet_view.rightToLeft=True
lines=[
 ('جدول مراجعة بنك الأسئلة — «الهندسة للأطفال»', True),
 ('لاختصاصي اللغة والتربية. المصدر: بنك الأسئلة بعد تصحيحات المرحلة 0 (20 أكتوبر 2026).', False),
 ('', False),
 ('ما المطلوب؟', True),
 ('1) راجع كل بند: الصحة اللغوية، والوضوح لعمر المرحلة، وتوحيد المصطلحات (دليل المصطلحات في الميثاق، القسم 7)، وقاعدة «فكّر معي» لا «أخطأت».', False),
 ('2) اكتب في الأعمدة الصفراء فقط: النص المعتمد، ونص النطق المشكول، والحالة، وملاحظاتك. لا تعدّل الأعمدة الرمادية، فالمعرّف يربط البند بالتطبيق.', False),
 ('3) إن كان النص الحالي سليماً فاترك «النص المعتمد» فارغاً واختر الحالة «معتمد كما هو».', False),
 ('4) العمود الوردي «ملاحظة التدقيق» فيه ملاحظات لغوية وتربوية من التدقيق الآلي لم تُطبَّق بعد؛ خذها اقتراحاً لا إلزاماً.', False),
 ('5) التشكيل: كامل في نص النطق لكل المراحل، وكامل في النص المعروض للمرحلة 1، وما يزيل اللبس في غيرها.', False),
 ('6) الأرقام تُكتب كما هي (التطبيق يحوّلها مشرقية أو غربية حسب اختيار الطفل). في نص النطق اكتب الأعداد الملتبسة كلمات: «٢٫٥» ← «اثنانِ فاصلةُ خمسة».', False),
 ('', False),
 ('الأوراق', True),
 ('ورقة لكل مرحلة. الصف الأول في كل ورقة رؤوس الأعمدة، وكل صف بعده بند واحد. ورقة «مثال» تُظهر الصف المملوء كما نتوقعه.', False),
 ('', False),
 ('ملخّص التقدّم (يُحدَّث تلقائياً من عمود الحالة)', True),
]
for i,(t,b) in enumerate(lines,1):
    c=L.cell(row=i,column=1,value=t); c.font=Font(name=FONT,bold=b,size=14 if i==1 else 11); c.alignment=Alignment(wrap_text=True,vertical='top')
L.column_dimensions['A'].width=110
r0=len(lines)+1
L.cell(row=r0,column=1,value='المرحلة').font=Font(name=FONT,bold=True); L.cell(row=r0,column=2,value='البنود').font=Font(name=FONT,bold=True); L.cell(row=r0,column=3,value='تمت مراجعتها').font=Font(name=FONT,bold=True)
L.column_dimensions['B'].width=12; L.column_dimensions['C'].width=16
heads=['المعرّف','المفهوم','اسم المفهوم','الوحدة','النوع','الشكل','السؤال (الحالي)','الإجابة الصحيحة','خيار خاطئ 1','خيار خاطئ 2','خيار خاطئ 3','التلميح','خطوات الحل','ملاحظة التدقيق',
       'السؤال المعتمد','الإجابة المعتمدة','الخيارات الخاطئة المعتمدة (افصل بـ ؛)','التلميح المعتمد','خطوات الحل المعتمدة (افصل بـ ؛)','نص النطق المشكول (للسؤال)','الحالة','ملاحظات الاختصاصي']
widths=[10,8,24,30,8,8,48,30,24,24,24,30,44,46,48,30,36,30,44,48,16,36]
INPUT_FROM=15
statuses=['لم يُراجَع','معتمد كما هو','معدَّل','يحتاج نقاشاً','يُحذف']
stage_rows={}
for st in '1234':
    ws=wb.create_sheet(stage_names[st]); ws.sheet_view.rightToLeft=True; ws.freeze_panes='C2'
    for j,h in enumerate(heads,1):
        c=ws.cell(row=1,column=j,value=h); c.font=Font(name=FONT,bold=True,color='FFFFFF'); c.fill=hdr_fill if j<INPUT_FROM else PatternFill('solid',fgColor='B8860B')
        c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center'); c.border=B
        ws.column_dimensions[get_column_letter(j)].width=widths[j-1]
    ws.row_dimensions[1].height=42
    dv=DataValidation(type='list',formula1='"'+','.join(statuses)+'"',allow_blank=False); ws.add_data_validation(dv)
    r=2
    for u in meta['units']:
        if u[1]!=st: continue
        for cid in meta['uc'].get(u,[]):
            if c2u.get(cid)!=u or cid not in bank: continue
            for x in bank[cid]:
                wr=(x.get('wrongs') or [])+['','','']
                vals=[x['id'],cid,meta['names'].get(cid,''),meta['titles'].get(u,u),x['kind'],'اختيار' if x['type']=='choice' else 'رقم',x['q'],
                      x.get('correct') if x['type']=='choice' else (str(x['ans'])+(f" (±{x['tol']})" if x.get('tol') else '')),wr[0],wr[1],wr[2],x.get('hint',''),' ← '.join(x.get('steps',[])),
                      '\n'.join(notes.get(x['id'],[])), None,None,None,None,None,None,'لم يُراجَع',None]
                for j,v in enumerate(vals,1):
                    c=ws.cell(row=r,column=j,value=v); c.font=Font(name=FONT,size=10); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=B
                    if j>=INPUT_FROM: c.fill=in_fill
                    elif j==14 and v: c.fill=note_fill
                    else: c.fill=PatternFill('solid',fgColor='F2F2F2')
                dv.add(ws.cell(row=r,column=21)); r+=1
    stage_rows[st]=(ws.title,r-1)
for k,st in enumerate('1234'):
    title,last=stage_rows[st]; rr=r0+1+k
    L.cell(row=rr,column=1,value=title).font=Font(name=FONT)
    L.cell(row=rr,column=2,value=f"=COUNTA('{title}'!A2:A{last})").font=Font(name=FONT)
    L.cell(row=rr,column=3,value=f"=COUNTA('{title}'!U2:U{last})-COUNTIF('{title}'!U2:U{last},\"لم يُراجَع\")").font=Font(name=FONT)
tr=r0+5
L.cell(row=tr,column=1,value='المجموع').font=Font(name=FONT,bold=True)
L.cell(row=tr,column=2,value=f'=SUM(B{r0+1}:B{r0+4})').font=Font(name=FONT,bold=True)
L.cell(row=tr,column=3,value=f'=SUM(C{r0+1}:C{r0+4})').font=Font(name=FONT,bold=True)
# ورقة مثال
E=wb.create_sheet('مثال'); E.sheet_view.rightToLeft=True
ex=['C134-04','C134','الحادة والمنفرجة','المرحلة 2 · الوحدة 3: زوايا الأركان','فهم','اختيار','فوضى تقول: الزاوية التي ضلعاها طويلان منفرجة، والتي ضلعاها قصيران حادة. ما رأيك؟','خطأ؛ نوع الزاوية يحدده انفتاح الضلعين لا طولهما','صحيح دائماً','صحيح إذا كان الضلعان أطول من 10 سم','','لو مددنا ضلعي زاوية، هل يتغيّر انفتاحهما؟','إطالة الضلعين لا تغيّر مقدار الدوران ← فالزاوية الحادة تبقى حادة','',
    'فوضى تقول: الزاوية ذات الضلعين الطويلين منفرجة، وذات الضلعين القصيرين حادة. ما رأيك؟','','','','','فَوْضَى تَقُولُ: الزَّاوِيَةُ ذَاتُ الضِّلْعَيْنِ الطَّوِيلَيْنِ مُنْفَرِجَةٌ، وَذَاتُ الضِّلْعَيْنِ القَصِيرَيْنِ حَادَّةٌ. مَا رَأْيُكَ؟','معدَّل','اختصرتُ «التي ضلعاها» لسهولة النطق.']
for j,h in enumerate(heads,1):
    c=E.cell(row=1,column=j,value=h); c.font=Font(name=FONT,bold=True,color='FFFFFF'); c.fill=hdr_fill if j<INPUT_FROM else PatternFill('solid',fgColor='B8860B'); c.alignment=Alignment(wrap_text=True); E.column_dimensions[get_column_letter(j)].width=widths[j-1]
    c=E.cell(row=2,column=j,value=ex[j-1]); c.font=Font(name=FONT,size=10); c.alignment=Alignment(wrap_text=True,vertical='top'); c.fill=in_fill if j>=INPUT_FROM else PatternFill('solid',fgColor='F2F2F2')

# ورقة نصوص «لماذا؟» وأسماء المفاهيم
import importlib.util
spec=importlib.util.spec_from_file_location('c','/home/claude/bank/corrections_p0.py'); cm=importlib.util.module_from_spec(spec); spec.loader.exec_module(cm)
wnotes={}
for sev,c in ar['a1']+ar['a2']:
    m1=re.match(r'(C\d{3})',c[0]); txt=c[1]
    cid = m1.group(1) if m1 and not c[0].startswith('CONC') else (re.search(r'C\d{3}',c[0]).group(0) if c[0].startswith('CONC') else None)
    if c[0].startswith('CONC') or txt.startswith('CONCEPT_'):
        if cid and cid not in cm.WHY: wnotes.setdefault(cid,[]).append(f"[{sev}] {c[2]} ← المقترح: {c[3]}")
W=wb.create_sheet('نصوص لماذا'); W.sheet_view.rightToLeft=True; W.freeze_panes='B2'
wh=['المفهوم','اسم المفهوم','نص «لماذا؟» الحالي','ملاحظة التدقيق','اسم المفهوم المعتمد','نص «لماذا؟» المعتمد','نص النطق المشكول','الحالة','ملاحظات الاختصاصي']
ww=[8,28,60,50,28,60,60,16,36]
for j,h in enumerate(wh,1):
    c=W.cell(row=1,column=j,value=h); c.font=Font(name=FONT,bold=True,color='FFFFFF'); c.fill=hdr_fill if j<5 else PatternFill('solid',fgColor='B8860B'); c.alignment=Alignment(wrap_text=True,horizontal='center'); W.column_dimensions[get_column_letter(j)].width=ww[j-1]
dv2=DataValidation(type='list',formula1='"'+','.join(statuses)+'"'); W.add_data_validation(dv2)
r=2
for u in meta['units']:
    for cid in meta['uc'].get(u,[]):
        if c2u.get(cid)!=u: continue
        vals=[cid,meta['names'].get(cid,''),meta['why'].get(cid,''),'\n'.join(wnotes.get(cid,[])),None,None,None,'لم يُراجَع',None]
        for j,v in enumerate(vals,1):
            c=W.cell(row=r,column=j,value=v); c.font=Font(name=FONT,size=10); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=B
            c.fill=in_fill if j>=5 else (note_fill if (j==4 and v) else PatternFill('solid',fgColor='F2F2F2'))
        dv2.add(W.cell(row=r,column=8)); r+=1
print('why rows',r-2,'why notes',sum(len(v) for v in wnotes.values()))

wb.save('/home/claude/bank/جدول_مراجعة_البنك_للاختصاصي.xlsx')
print('notes attached',sum(len(v) for v in notes.values()),'unmatched',len(unmatched)); print(unmatched[:15])
