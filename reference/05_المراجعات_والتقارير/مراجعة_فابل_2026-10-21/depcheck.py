import openpyxl, sys, collections
p='/home/claude/out/حزمة_المرجع_الجديدة/04_السجلات/REGISTRY_v2_سجل_المفاهيم_الحي.xlsx'
wb=openpyxl.load_workbook(p, data_only=True)
C=wb['المفاهيم']; P=wb['المواضع']; D=wb['الاعتماديات']
concepts={}
for r in C.iter_rows(min_row=2, values_only=True):
    if r[0]: concepts[r[0]]={'name':r[1],'stage':r[2],'status':r[3]}
print('concepts',len(concepts))
# stage/unit/lesson ordering: stage-unit order. M3 units start at 0, M2 at 0.
positions=[]
for r in P.iter_rows(min_row=2, values_only=True):
    if r[0]: positions.append(dict(pid=r[0],cid=r[1],name=r[2],stage=r[3],unit=r[4],lesson=r[5],level=r[6],kind=r[7],pos=r[8],status=r[9]))
print('positions',len(positions))
levels=collections.Counter(p['level'] for p in positions); print(levels)
kinds=collections.Counter(p['kind'] for p in positions); print(kinds)
# key
def key(p): return (p['stage'],p['unit'],p['lesson'])
bycid=collections.defaultdict(list)
for p in positions: bycid[p['cid']].append(p)
# (c) concepts without position
nopos=[c for c in concepts if c not in bycid]
print('\n(c) concepts with no position:',len(nopos))
for c in nopos: print(' ',c,concepts[c])
# (d) positions whose concept not in sheet
unk=[p for p in positions if p['cid'] not in concepts]
print('\n(d) positions with unknown concept:',len(unk))
for p in unk: print(' ',p)
# first تأسيس position
first_found={}
first_any={}
for c,ps in bycid.items():
    f=[p for p in ps if p['level']=='تأسيس']
    if f: first_found[c]=min(key(p) for p in f)
    first_any[c]=min(key(p) for p in ps)
nofound=[c for c in bycid if c not in first_found]
print('\nconcepts with positions but no تأسيس:',len(nofound))
for c in nofound: print(' ',c,concepts.get(c,{}).get('name'),[ (p['level'],key(p)) for p in bycid[c]])
# dependencies
deps=[]
for r in D.iter_rows(min_row=2, values_only=True):
    if r[0] and r[2]: deps.append((r[0],r[2],r[4]))
print('\ndeps',len(deps))
# unknown ids in deps
for a,b,j in deps:
    if a not in concepts: print(' dep unknown concept',a)
    if b not in concepts: print(' dep unknown prereq',b)
# (a) cycles
g=collections.defaultdict(set)
for a,b,j in deps: g[a].add(b)
color={}
cycles=[]
def dfs(u,stack):
    color[u]=1; stack.append(u)
    for v in g[u]:
        if color.get(v,0)==0: dfs(v,stack)
        elif color.get(v)==1:
            cycles.append(stack[stack.index(v):]+[v])
    stack.pop(); color[u]=2
for u in list(g):
    if color.get(u,0)==0: dfs(u,[])
print('\n(a) cycles:',len(cycles))
for cyc in cycles: print(' ',' -> '.join(cyc))
# self deps
print(' self-deps:',[(a,b) for a,b,j in deps if a==b])
# duplicates
cnt=collections.Counter((a,b) for a,b,j in deps)
print(' duplicate dep rows:',[(k,v) for k,v in cnt.items() if v>1])
# (b) concept used before its first تأسيس: for each dep a->b, every position of a (any level? use the first position of a of level != 'استدعاء') must be >= first تأسيس of b
print('\n(b) prerequisite not founded before use:')
viol=[]
for a,b,j in deps:
    if a not in bycid or b not in concepts: continue
    use=min(key(p) for p in bycid[a] if p['level'] in ('تأسيس','مدخل','إتقان'))
    fb=first_found.get(b)
    if fb is None:
        viol.append((a,b,use,None,j,'المتطلب بلا تأسيس'))
    elif fb>use:
        viol.append((a,b,use,fb,j,'المتطلب يُؤسَّس بعد الاستعمال'))
    elif fb==use:
        viol.append((a,b,use,fb,j,'الدرس نفسه'))
for v in sorted(viol,key=lambda x:(x[5],x[2])):
    print(' ',v, concepts[v[0]]['name'],'<-',concepts[v[1]]['name'])
# also check: first مدخل of a before first تأسيس of b where a has مدخل — included above
# positions with multiple تأسيس
print('\nconcepts with more than one تأسيس position:')
for c,ps in bycid.items():
    f=[key(p) for p in ps if p['level']=='تأسيس']
    if len(f)>1: print(' ',c,concepts[c]['name'],f)
# level order check: إتقان before تأسيس, استدعاء before تأسيس
print('\nlevel order anomalies:')
for c,ps in bycid.items():
    fb=first_found.get(c)
    for p in ps:
        if fb and p['level'] in ('استدعاء','إتقان') and key(p)<fb:
            print(' ',c,concepts[c]['name'],p['level'],key(p),'before تأسيس',fb)
# pos code consistency
print('\npos code mismatch:')
for p in positions:
    code=p['stage']*10000+p['unit']*100+p['lesson']
    if p['pos']!=code: print(' ',p['pid'],p['cid'],p['pos'],code)
# stage of concept vs stage of first position
print('\nstage mismatch (concept stage vs first تأسيس stage):')
for c,fb in first_found.items():
    if concepts[c]['stage']!=fb[0]: print(' ',c,concepts[c]['name'],'sheet stage',concepts[c]['stage'],'first تأسيس',fb)
# judgments column distribution
print('\njudgments:',collections.Counter(j for a,b,j in deps))
