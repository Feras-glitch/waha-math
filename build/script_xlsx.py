# يكتب جدول التسجيل xlsx من الصفوف التي يولّدها build/script.mjs (التنسيق نفسه الذي كان في extract_script.py).
# الاستعمال (من البناء): python3 -I build/script_xlsx.py dist/script_m1.json dist/script_m1.xlsx
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
SRC, OUT = sys.argv[1], sys.argv[2]
d = json.load(open(SRC, encoding='utf-8'))
wb = Workbook(); ws = wb.active; ws.title = 'جدول التسجيل'; ws.sheet_view.rightToLeft = True
ws.append(['#', 'الدرس', 'المشهد', 'الشخصية', 'المعرّف', 'النص للولد', 'النص للبنت', 'ملف الولد', 'ملف البنت', 'توجيه الأداء', 'ملاحظة الاختصاصي', 'الحالة'])
for r in d['rows']: ws.append(r)
hdr = PatternFill('solid', fgColor='1F5F99'); thin = Side(style='thin', color='CFE0E0')
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True); c.border = Border(bottom=thin)
    if row[3].value == 'فوضى':
        for c in row: c.fill = PatternFill('solid', fgColor='FDECEA')
for col, w in zip('ABCDEFGHIJKL', [5, 14, 16, 9, 30, 48, 48, 32, 32, 30, 30, 12]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'F2'; ws.auto_filter.ref = ws.dimensions
s2 = wb.create_sheet('ملخص'); s2.sheet_view.rightToLeft = True
s2.append(['الدرس', 'الشخصية', 'عدد الجمل'])
for l, w, n in d['summary']: s2.append([l, w, n])
s2.append([]); s2.append(['مجموع الجمل', '', d['nclips']]); s2.append(['مجموع الملفات الصوتية (مع صيغتي الولد والبنت)', '', d['nfiles']])
s2.append([]); s2.append(['قواعد التسمية', '', 'اسم الملف = المعرّف حرفاً بحرف + .mp3. الجملة ذات الصيغتين لها ملفان ‎_m‎ و‎_f‎. الجمل العامة (gen) تُسجَّل مرة واحدة لكل الدروس.'])
s2.column_dimensions['A'].width = 40; s2.column_dimensions['C'].width = 80
for c in s2[1]: c.font = Font(bold=True)
# بناء حتمي: الملف نفسه بايتاً بايتاً للمحتوى نفسه (لا تاريخ إنشاء ولا أوقات في أرشيف zip)، فيُقارَن الجدول بين بناءين
import datetime, io, zipfile
FIXED = datetime.datetime(2026, 1, 1)
wb.properties.created = FIXED; wb.properties.modified = FIXED; wb.properties.lastModifiedBy = None; wb.properties.creator = 'واحة الرياضيات'
buf = io.BytesIO(); wb.save(buf); src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as z:
    for info in src.infolist():
        data = src.read(info.filename)
        if info.filename == 'docProps/core.xml':   # openpyxl يكتب وقت الحفظ في dcterms:modified مهما ضُبط
            import re; data = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(<)', rb'\g<1>2026-01-01T00:00:00Z\2', data)
        zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0)); zi.compress_type = zipfile.ZIP_DEFLATED; zi.external_attr = 0o644 << 16
        z.writestr(zi, data)
