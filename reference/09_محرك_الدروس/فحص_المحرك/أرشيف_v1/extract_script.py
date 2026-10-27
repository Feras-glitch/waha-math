# يستخرج كل جملة منطوقة من ملف الوحدة إلى جدول تسجيل للاختصاصي والمدبلجين.
# كل جملة نصٌّ ثابت بمعرّف ثابت؛ الجمل العامة gen.* مشتركة، والباقي يُسبق بمعرّف الدرس.
import re, sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
SRC, OUT = sys.argv[1], sys.argv[2]
s = open(SRC, encoding='utf-8').read()
js = s[s.index('<script>'):]
# حدود المناطق: المحرّك قبل الدرس 1، ثم الدروس، ثم مشغّل الدرس (محرّك)
marks = [(m.start(), m.group(1)) for m in re.finditer(r'/\* -+ \[(m\d\.u\d\.d\d)\]', js)]
end = js.index('مشغّل الدرس')
def region(pos):
    if pos > end: return None
    cur = None
    for p, lid in marks:
        if pos >= p: cur = lid
    return cur
AR = re.compile(r'[؀-ۿ]')
rows = []
def add(pos, who, lid, text):
    rows.append(dict(pos=pos, who=who, id=lid, text=text))
for m in re.finditer(r"say\('(compass|chaos)', '([a-z][\w.]*)', '([^']*)'(?:\)|, \{)", js): add(m.start(), m.group(1), m.group(2), m.group(3))
for m in re.finditer(r"\['(compass|chaos)', '([a-z][\w.]*)', '([^']*)'\]", js): add(m.start(), m.group(1), m.group(2), m.group(3))
for m in re.finditer(r"\['([a-z]\w*\.[\w.]*)', '([^']*)'\]", js):   # معرّف الجملة فيه نقطة دائماً؛ أزرار الاختيار ليست جملاً منطوقة
    if AR.search(m.group(2)): add(m.start(), 'compass', m.group(1), m.group(2))
for m in re.finditer(r"id: '([a-z]\w*\.[\w.]*)', t: '([^']*)'", js): add(m.start(), 'compass', m.group(1), m.group(2))   # أسباب «صحّح لفوضى» وخيارات «لماذا؟» في التحدي تُقرأ بالصوت
problems = []; clips = {}
for r in rows:
    reg = region(r['pos'])
    if r['id'].startswith('gen.'): full = r['id']
    elif reg: full = reg + '.' + r['id']
    else: problems.append(('جملة خارج الدروس بلا gen.', r['id'], r['text'])); continue
    if '{name}' in r['text']: problems.append(('اسم الطالب في جملة منطوقة (لا يُسجَّل)', full, r['text']))
    if full in clips and clips[full]['text'] != r['text']: problems.append(('معرّف واحد لنصّين مختلفين', full, clips[full]['text'] + ' ⟂ ' + r['text']))
    clips.setdefault(full, dict(full=full, who=r['who'], text=r['text'], lesson=reg if not r['id'].startswith('gen.') else 'عام'))
GEN = re.compile(r'\{([^|{}]*)\|([^|{}]*)\}')
def form(t, i): return GEN.sub(lambda m: m.group(1 + i), t)
SCENE = [('hook', 'إشعال'), ('exp', 'اكتشاف'), ('wk', 'مثال محلول'), ('pr', 'تمرين'), ('why', 'تمرين: سبب الصواب'), ('gen.correct', 'تمرين: مدح'), ('gen.err', 'تمرين: تفسير الخطأ'),
         ('gen.solve', 'تمرين: حل مشترك'), ('gen.twin', 'تمرين: مسألة توأم'), ('gen.hint', 'مساعدة'), ('gen.demo', 'مساعدة'), ('gen.num', 'عدّ'), ('chaos', 'صحّح لفوضى'), ('gen.chaos', 'صحّح لفوضى'),
         ('ms', 'فحص الإتقان'), ('gen.ms', 'فحص الإتقان'), ('wrap', 'ختام'), ('rv', 'مراجعة'), ('test', 'التحدي'), ('ts', 'التحدي'), ('gen.hintOffer', 'مساعدة')]
def scene(full):
    tail = full.split('.', 3)[-1] if full.startswith('m1.') else full
    best = ''
    for k, v in SCENE:
        if tail.startswith(k) and len(k) > len(best): best, lab = k, v
    return lab if best else '؟'
NOTE = {'إشعال': 'تشويق ودهشة، إيقاع حيّ', 'اكتشاف': 'دعوة هادئة للتجربة، مهل', 'مثال محلول': 'شرح واضح بطيء، وقفة عند كل عدّ',
        'تمرين': 'تشجيع', 'تمرين: مدح': 'فرح صادق قصير', 'تمرين: سبب الصواب': 'تأكيد هادئ للفكرة', 'تمرين: تفسير الخطأ': 'لطيف غير ناقد؛ «لنفكّر معاً» بنبرة شراكة',
        'تمرين: حل مشترك': 'دافئ، بلا خيبة', 'تمرين: مسألة توأم': 'تحفيز', 'مساعدة': 'همس ودود', 'عدّ': 'عدّ واضح منغّم', 'صحّح لفوضى': 'فوضى: مشاكسة مرحة بلا عدوانية؛ فرجار: فضول',
        'فحص الإتقان': 'جِدّ لطيف بلا ضغط', 'ختام': 'احتفال وفخر', 'مراجعة': 'تذكير واضح', 'التحدي': 'سؤال واضح محايد، بلا تلميح بالنبرة'}
WHO = {'compass': 'فِرجار', 'chaos': 'فوضى'}
out = sorted(clips.values(), key=lambda c: (c['lesson'] == 'عام', c['lesson'], list(clips).index(c['full'])))
wb = Workbook(); ws = wb.active; ws.title = 'جدول التسجيل'; ws.sheet_view.rightToLeft = True
H = ['#', 'الدرس', 'المشهد', 'الشخصية', 'المعرّف', 'النص للولد', 'النص للبنت', 'ملف الولد', 'ملف البنت', 'توجيه الأداء', 'ملاحظة الاختصاصي', 'الحالة']
ws.append(H)
nfiles = 0
for i, c in enumerate(out, 1):
    two = bool(GEN.search(c['text']))
    m, f = form(c['text'], 0), form(c['text'], 1)
    fm = c['full'] + ('_m' if two else '') + '.mp3'; ff = c['full'] + '_f.mp3' if two else '(الملف نفسه)'
    nfiles += 2 if two else 1
    sc = scene(c['full'])
    ws.append([i, c['lesson'].replace('m1.u', 'الوحدة ').replace('.d', '، الدرس ') if c['lesson'] != 'عام' else 'عام (كل الدروس)', sc, WHO[c['who']], c['full'], m, f if two else '(النص نفسه)', fm, ff, NOTE.get(sc, ''), '', ''])
hdr = PatternFill('solid', fgColor='1F5F99'); thin = Side(style='thin', color='CFE0E0')
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True); c.border = Border(bottom=thin)
    if row[3].value == 'فوضى':
        for c in row: c.fill = PatternFill('solid', fgColor='FDECEA')
for col, w in zip('ABCDEFGHIJKL', [5, 14, 16, 9, 30, 48, 48, 32, 32, 30, 30, 12]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'F2'; ws.auto_filter.ref = ws.dimensions
s2 = wb.create_sheet('ملخص'); s2.sheet_view.rightToLeft = True
cnt = {}
for c in out: k = (c['lesson'], WHO[c['who']]); cnt[k] = cnt.get(k, 0) + 1
s2.append(['الدرس', 'الشخصية', 'عدد الجمل'])
for (l, w), n in sorted(cnt.items()): s2.append([l, w, n])
s2.append([]); s2.append(['مجموع الجمل', '', len(out)]); s2.append(['مجموع الملفات الصوتية (مع صيغتي الولد والبنت)', '', nfiles])
s2.append([]); s2.append(['قواعد التسمية', '', 'اسم الملف = المعرّف حرفاً بحرف + .mp3. الجملة ذات الصيغتين لها ملفان ‎_m‎ و‎_f‎. الجمل العامة (gen) تُسجَّل مرة واحدة لكل الدروس.'])
s2.column_dimensions['A'].width = 40; s2.column_dimensions['C'].width = 80
for c in s2[1]: c.font = Font(bold=True)
wb.save(OUT)
json.dump(sorted({c['full'] + ('_m' if GEN.search(c['text']) else '') for c in out} | {c['full'] + '_f' for c in out if GEN.search(c['text'])}), open(OUT + '.ids.json', 'w'), ensure_ascii=False)
print('جمل:', len(out), '| ملفات:', nfiles, '| مشاكل:', len(problems))
for p in problems: print(' -', *p)
